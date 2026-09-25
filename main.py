import os
import contextlib
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import yfinance as yf
import pandas as pd
from datetime import datetime

# Import our custom modules
from macro_data import get_macro_inputs, prompt_float
from beta_comps import run_beta_comps

# --- THEME & STYLING CONSTANTS ---
ALIGN_LEFT = Alignment(horizontal='left')
ALIGN_RIGHT = Alignment(horizontal='right')
ALIGN_CENTER = Alignment(horizontal='center')

FONT_MAIN = Font(name="Calibri", size=11, color="000000")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="000000")
FONT_ITALIC = Font(name="Calibri", size=11, italic=True, color="000000")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="002060")  # Dark Blue Headers

# Wall Street Standard: Blue = hardcoded inputs, Black = formulas, Green = links to other sheets
FONT_INPUT = Font(name="Calibri", size=11, color="0000FF")
FONT_FORMULA = Font(name="Calibri", size=11, color="000000")
FONT_LINK = Font(name="Calibri", size=11, color="008000")
FONT_WARN = Font(name="Calibri", size=11, bold=True, color="FF0000")

FILL_HEADER = PatternFill(start_color="DCE6F1", end_color="DCE6F1", fill_type="solid")  # Light Blue
FILL_INPUT = PatternFill(start_color="FFFFE0", end_color="FFFFE0", fill_type="solid")   # Light Yellow
FILL_RESULT = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")  # Light Gray for outputs

BORDER_BOTTOM = Border(bottom=Side(style='thin'))
BORDER_TB = Border(top=Side(style='thin'), bottom=Side(style='thin'))
BORDER_DOTTED = Border(bottom=Side(style='dotted'))  # For subsections
BORDER_TOP_BOTTOM_THICK = Border(top=Side(style='thin'), bottom=Side(style='medium'))

# Synthetic-rating lookup: (minimum EBIT / interest expense, rating, default spread over Rf).
# APPROXIMATE, modelled on Damodaran's large-cap table. Refresh these numbers from his current
# "ratings / interest coverage / default spread" dataset before you cite them anywhere.
SYNTHETIC_SPREADS = [
    (8.50, "AAA", 0.0060), (6.50, "AA", 0.0070), (5.50, "A+", 0.0090), (4.25, "A", 0.0100),
    (3.00, "A-", 0.0110), (2.50, "BBB", 0.0140), (2.25, "BB+", 0.0200), (2.00, "BB", 0.0250),
    (1.75, "B+", 0.0330), (1.50, "B", 0.0400), (1.25, "B-", 0.0500), (0.80, "CCC", 0.0700),
    (0.65, "CC", 0.0950), (0.20, "C", 0.1150), (float("-inf"), "D", 0.1400),
]


def apply_style(cell, font=FONT_MAIN, fill=None, num_format=None, align=None, border=None):
    cell.font = font
    if fill: cell.fill = fill
    if num_format: cell.number_format = num_format
    if align: cell.alignment = align
    if border: cell.border = border


# ----------------------------------------------------------------------------------------------
# DATA HELPERS
# ----------------------------------------------------------------------------------------------
def get_fx_rate(from_ccy, to_ccy):
    """Latest spot FX (units of to_ccy per 1 from_ccy) from Yahoo, or None if unavailable."""
    if from_ccy == to_ccy:
        return 1.0
    try:
        with contextlib.redirect_stderr(io.StringIO()):
            hist = yf.Ticker(f"{from_ccy}{to_ccy}=X").history(period="5d")
        if not hist.empty:
            return float(hist["Close"].dropna().iloc[-1])
    except Exception:
        pass
    return None


def get_live_financials(ticker):
    """
    Pulls live market cap and total debt, both expressed in the stock's TRADING currency.
    Yahoo reports market cap in the trading currency but balance-sheet items in the reporting
    currency (e.g. Air China: HKD market cap, RMB debt), so debt is converted when they differ.
    """
    out = {"debt": 0.0, "mcap": 0.0, "ccy": None, "ok": False, "note": ""}
    try:
        t = yf.Ticker(ticker)
        info = t.info or {}
        mcap = float(info.get("marketCap") or 0)
        debt = float(info.get("totalDebt") or 0)

        if not debt:
            bs = t.balance_sheet
            if not bs.empty and "Total Debt" in bs.index:
                debt = float(bs.loc["Total Debt"].iloc[0])

        ccy = info.get("currency")
        fin_ccy = info.get("financialCurrency") or ccy
        if ccy and fin_ccy and ccy != fin_ccy:
            fx = get_fx_rate(fin_ccy, ccy)
            if fx is None:
                out["note"] = f"{ticker}: debt in {fin_ccy} but market cap in {ccy}; FX lookup failed, NOT converted"
            else:
                debt *= fx
                out["note"] = f"{ticker}: debt converted {fin_ccy}->{ccy} at {fx:.4f}"

        out.update(debt=debt, mcap=mcap, ccy=ccy, ok=(mcap > 0))
    except Exception as e:
        out["note"] = f"{ticker}: financials pull failed ({e})"
    return out


def _latest(df, labels):
    """First non-null latest-period value among candidate row labels of a yfinance statement."""
    if df is None or getattr(df, "empty", True):
        return None
    for lab in labels:
        if lab in df.index:
            vals = df.loc[lab].dropna()
            if len(vals):
                return float(vals.iloc[0])
    return None


def synthetic_rating(icr):
    for floor, rating, spread in SYNTHETIC_SPREADS:
        if icr >= floor:
            return rating, spread
    return SYNTHETIC_SPREADS[-1][1], SYNTHETIC_SPREADS[-1][2]


def get_cost_of_debt_spec(macro):
    """
    Returns a dict describing how the pre-tax cost of debt is built:
      {"mode": "spread",     "spread": x, "note": str}   Kd = Rf + credit spread   (default)
      {"mode": "historical", "interest": x, "debt": y}   Kd = interest / debt
      {"mode": "manual",     "kd": x}
    """
    print("\n--- COST OF DEBT (Kd) ---")
    print("  [1] Risk-free rate + credit spread (synthetic rating from interest coverage)  [default]")
    print("  [2] Historical: interest expense / total debt (average coupon, not a marginal cost)")
    print("  [3] Enter manually (from a rating or bond yield)")
    choice = input("Select method [1]: ").strip()

    inc = bs = None
    if choice in ("", "1", "2"):
        try:
            print("Fetching Target's income statement and balance sheet...")
            t = yf.Ticker(macro["ticker"])
            inc, bs = t.financials, t.balance_sheet
        except Exception:
            print("❌ Financials pull failed.")

    int_exp = _latest(inc, ["Interest Expense", "Interest Expense Non Operating"])
    int_exp = abs(int_exp) if int_exp else None
    ebit = _latest(inc, ["EBIT", "Operating Income"])
    tot_debt = _latest(bs, ["Total Debt"])

    if choice == "2":
        if int_exp and tot_debt:
            print(f"✓ Interest expense {int_exp:,.0f} / total debt {tot_debt:,.0f} = {int_exp / tot_debt:.2%}")
            if int_exp / tot_debt <= macro["risk_free_rate"]:
                print("  [!] That is at or below the risk-free rate, so it cannot be a marginal borrowing cost.")
                print("      The sheet will flag it. Consider method 1.")
            return {"mode": "historical", "interest": int_exp, "debt": tot_debt}
        print("❌ Could not locate interest expense / total debt. Falling back to manual entry.")
        choice = "3"

    if choice in ("", "1"):
        suggested, note = None, ""
        if ebit is not None and int_exp:
            icr = ebit / int_exp
            rating, suggested = synthetic_rating(icr)
            note = f"Synthetic rating {rating} (EBIT / interest = {icr:.1f}x); spread table approximate"
            print(f"✓ EBIT / interest expense = {icr:.1f}x -> synthetic rating {rating}, approx. spread {suggested:.2%}")
        else:
            print("❌ Could not compute interest coverage, so no spread can be suggested.")
        hint = f"[Default {suggested * 100:.2f}%]" if suggested is not None else "[required]"
        spread = prompt_float(f"Credit spread over risk-free % {hint}: ",
                              suggested * 100 if suggested is not None else None, "credit spread")
        if not note or abs(spread - (suggested or 0)) > 1e-9:
            note = "Credit spread entered by user"
        return {"mode": "spread", "spread": spread, "note": note}

    kd_val = prompt_float("Enter Pre-tax Cost of Debt % (from credit rating/bonds): ", None, "cost of debt")
    return {"mode": "manual", "kd": kd_val}


def get_erp_spec():
    """ERP is a sourced input. Historical estimates stay on the COST OF EQUITY sheet for reference only."""
    print("\n--- EQUITY RISK PREMIUM ---")
    print("  [1] Enter a sourced ERP (recommended: e.g. implied ERP for the market + country risk premium)  [default]")
    print("  [2] Fall back to the historical geometric estimate (unstable, reference only)")
    choice = input("Select method [1]: ").strip()
    if choice == "2":
        return {"mode": "historical",
                "source": "FALLBACK: historical geometric return + dividend yield - Rf (unstable; replace with a sourced ERP)"}
    erp = prompt_float("Enter Equity Risk Premium %: ", None, "equity risk premium")
    src = input("Source / citation for this ERP (e.g. 'Damodaran, Jan-2026, India'): ").strip() or "User input (source not stated)"
    return {"mode": "input", "erp": erp, "source": src}


# ----------------------------------------------------------------------------------------------
# SHEET BUILDERS
# ----------------------------------------------------------------------------------------------
def build_wacc_sheet(ws, macro, comps_data, beta_settings, kd_spec, erp_spec):
    """
    Builds the main WACC dashboard with strict IB formatting.
    Returns dict of key cell rows so other sheets can link back to it.
    """
    ws.title = "WACC"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions['A'].width = 1.15

    ws["B1"] = "Weighted Average Cost of Capital"
    apply_style(ws["B1"], font=Font(name="Calibri", size=14, bold=True, color="002060"), fill=FILL_HEADER)
    ws.merge_cells("B1:N1")

    ws["B3"] = "Debt and equity are in each company's own trading currency (INR in crores, all other currencies in billions)."
    apply_style(ws["B3"], font=FONT_ITALIC)

    # --- PEER COMPS SECTION ---
    ws["B5"] = "Peer Comps"
    apply_style(ws["B5"], font=FONT_BOLD)

    for col in range(2, 15):
        ws.cell(row=5, column=col).border = BORDER_BOTTOM

    headers = ["Name", "Country", "Total Debt", "Total Equity", "Tax Rate\u00B9", "Debt/Equity", "Debt/Capital",
               "Levered Beta\u00B2", "Unlevered Beta\u00B3", "Obs", "R\u00B2", "Currency", "Raw Beta 95% CI"]
    for col_idx, header in enumerate(headers, 2):
        cell = ws.cell(row=7, column=col_idx, value=header)
        apply_style(cell, font=FONT_HEADER, border=BORDER_TB)

    row = 8
    target_row = None
    peer_rows = []

    for peer in comps_data:
        is_target = peer['Role'] == 'Target'
        if is_target:
            target_row = row
        else:
            peer_rows.append(row)

        fin = get_live_financials(peer['Ticker'])
        if fin["note"]:
            print(f"  [i] {fin['note']}")
        if not fin["ok"]:
            print(f"  [!] {peer['Ticker']}: market cap unavailable, row will show n/a and is ignored in peer statistics.")
        scale = 1e7 if fin["ccy"] == "INR" else 1e9
        t = peer['Ticker']

        ws.cell(row=row, column=2, value=f"{t} (Target)" if is_target else t).font = FONT_BOLD if is_target else FONT_MAIN
        ws.cell(row=row, column=3, value=peer['Country']).font = FONT_MAIN

        c_debt = ws.cell(row=row, column=4, value=fin["debt"] / scale)
        c_eq = ws.cell(row=row, column=5, value=fin["mcap"] / scale)
        c_tax = ws.cell(row=row, column=6, value=peer['Tax Rate'])
        c_lb = ws.cell(row=row, column=9, value=f"='Beta - {t}'!$K$11")

        apply_style(c_debt, font=FONT_INPUT, num_format='#,##0.0')
        apply_style(c_eq, font=FONT_INPUT, num_format='#,##0.0')
        apply_style(c_tax, font=FONT_INPUT, num_format='0.00%')
        apply_style(c_lb, font=FONT_LINK, num_format='0.00')

        c_de = ws.cell(row=row, column=7, value=f'=IFERROR(D{row}/E{row},"n/a")')
        c_dc = ws.cell(row=row, column=8, value=f'=IFERROR(D{row}/(D{row}+E{row}),"n/a")')
        c_ub = ws.cell(row=row, column=10, value=f'=IFERROR(I{row}/(1+(1-F{row})*G{row}),"n/a")')
        for c in [c_de, c_dc]: apply_style(c, font=FONT_FORMULA, num_format='0.00%')
        apply_style(c_ub, font=FONT_FORMULA, num_format='0.00')

        c_obs = ws.cell(row=row, column=11, value=f"='Beta - {t}'!$K$20")
        c_r2 = ws.cell(row=row, column=12, value=f"='Beta - {t}'!$K$18")
        c_ccy = ws.cell(row=row, column=13, value=fin["ccy"] or "n/a")
        c_ci = ws.cell(row=row, column=14,
                       value=f"=IFERROR(TEXT('Beta - {t}'!$K$22,\"0.00\")&\" to \"&TEXT('Beta - {t}'!$K$23,\"0.00\"),\"n/a\")")
        apply_style(c_obs, font=FONT_LINK, num_format='0')
        apply_style(c_r2, font=FONT_LINK, num_format='0.00')
        apply_style(c_ccy, font=FONT_INPUT)
        apply_style(c_ci, font=FONT_LINK, align=ALIGN_RIGHT)

        if is_target:
            for col in range(2, 15):
                ws.cell(row=row, column=col).border = BORDER_DOTTED

        row += 1

    # Peer statistics EXCLUDE the target. Fall back to the whole table only if there are no peers.
    if peer_rows:
        stat_start, stat_end = min(peer_rows), max(peer_rows)
    else:
        print("  [!] No peers available; statistics fall back to the target row alone.")
        stat_start = stat_end = target_row
    peer_end_row = row - 1  # last table row (used for stat row offsets below)

    c_avg_label = ws.cell(row=row + 1, column=5, value="Peer Average")
    apply_style(c_avg_label, font=FONT_BOLD, border=BORDER_TB)

    c_med_label = ws.cell(row=row + 2, column=5, value="Peer Median")
    apply_style(c_med_label, font=FONT_BOLD, border=BORDER_BOTTOM)

    for c in range(6, 11):
        col_let = openpyxl.utils.get_column_letter(c)
        c_avg = ws.cell(row=row + 1, column=c, value=f"=AVERAGE({col_let}{stat_start}:{col_let}{stat_end})")
        c_med = ws.cell(row=row + 2, column=c, value=f"=MEDIAN({col_let}{stat_start}:{col_let}{stat_end})")

        fmt = '0.00%' if c < 9 else '0.00'
        apply_style(c_avg, font=FONT_FORMULA, num_format=fmt, border=BORDER_TB)
        apply_style(c_med, font=FONT_FORMULA, num_format=fmt, border=BORDER_BOTTOM)

    avg_row, med_row = row + 1, row + 2

    # --- SUBSECTION: COST OF DEBT & EQUITY ---
    r_cde = row + 4
    ws.cell(row=r_cde, column=2, value="Cost of Debt").font = FONT_BOLD
    ws.cell(row=r_cde, column=7, value="Cost of Equity").font = FONT_BOLD

    for c in range(2, 5): ws.cell(row=r_cde, column=c).border = BORDER_DOTTED
    for c in range(7, 9): ws.cell(row=r_cde, column=c).border = BORDER_DOTTED

    rf_row = r_cde + 2
    kd_row = r_cde + 2
    ws.cell(row=kd_row, column=2, value="Pre-tax Cost of Debt").font = FONT_MAIN
    c_kd = ws.cell(row=kd_row, column=4)

    mode = kd_spec["mode"]
    if mode == "spread":
        ws.cell(row=r_cde + 5, column=2, value="  Credit spread over Rf").font = FONT_ITALIC
        c_sp = ws.cell(row=r_cde + 5, column=4, value=kd_spec["spread"])
        apply_style(c_sp, font=FONT_INPUT, fill=FILL_INPUT, num_format='0.00%')
        ws.cell(row=r_cde + 6, column=2, value=f"  {kd_spec['note']}").font = FONT_ITALIC
        c_kd.value = f"=H{rf_row}+D{r_cde + 5}"
        apply_style(c_kd, font=FONT_FORMULA, num_format='0.00%')
    elif mode == "historical":
        ws.cell(row=r_cde + 5, column=2, value="  Interest expense (latest FY)").font = FONT_ITALIC
        apply_style(ws.cell(row=r_cde + 5, column=4, value=kd_spec["interest"]), font=FONT_INPUT, num_format='#,##0')
        ws.cell(row=r_cde + 6, column=2, value="  Total debt (latest FY)").font = FONT_ITALIC
        apply_style(ws.cell(row=r_cde + 6, column=4, value=kd_spec["debt"]), font=FONT_INPUT, num_format='#,##0')
        c_kd.value = f"=D{r_cde + 5}/D{r_cde + 6}"
        apply_style(c_kd, font=FONT_FORMULA, num_format='0.00%')
    else:
        c_kd.value = kd_spec["kd"]
        apply_style(c_kd, font=FONT_INPUT, num_format='0.00%')

    # Sanity flag: a company cannot borrow below its own government's yield
    c_flag = ws.cell(row=kd_row, column=5, value=f'=IF(D{kd_row}<=H{rf_row},"<- Kd at/below Rf: check","")')
    apply_style(c_flag, font=FONT_WARN)

    ws.cell(row=r_cde + 3, column=2, value="Tax Rate").font = FONT_MAIN
    c_tax_target = ws.cell(row=r_cde + 3, column=4, value=macro['tax_rate'])
    apply_style(c_tax_target, font=FONT_INPUT, num_format='0.00%')

    ws.cell(row=r_cde + 4, column=2, value="After Tax Cost of Debt").font = FONT_MAIN
    c_kd_post = ws.cell(row=r_cde + 4, column=4, value=f"=D{kd_row}*(1-D{r_cde + 3})")
    apply_style(c_kd_post, font=FONT_FORMULA, num_format='0.00%', border=BORDER_TB)

    ws.cell(row=rf_row, column=7, value="Risk Free Rate").font = FONT_MAIN
    c_rf = ws.cell(row=rf_row, column=8, value=macro['risk_free_rate'])
    apply_style(c_rf, font=FONT_INPUT, num_format='0.00%')

    ws.cell(row=r_cde + 3, column=7, value="Equity Risk Premium").font = FONT_MAIN
    c_erp = ws.cell(row=r_cde + 3, column=8, value="='COST OF EQUITY'!$G$17")
    apply_style(c_erp, font=FONT_LINK, num_format='0.00%')

    r_cap = r_cde + 8
    ws.cell(row=r_cde + 4, column=7, value="Levered Beta").font = FONT_MAIN
    c_relevered = ws.cell(row=r_cde + 4, column=8, value=f"=H{r_cap + 5}")
    apply_style(c_relevered, font=FONT_FORMULA, num_format='0.00')

    ws.cell(row=r_cde + 5, column=7, value="Cost of Equity").font = FONT_MAIN
    c_ke = ws.cell(row=r_cde + 5, column=8, value=f"=H{rf_row}+(H{r_cde + 4}*H{r_cde + 3})")
    apply_style(c_ke, font=FONT_FORMULA, num_format='0.00%', border=BORDER_TB)

    # --- SUBSECTION: CAPITAL STRUCTURE & RELEVERING ---
    ws.cell(row=r_cap, column=2, value="Capital Structure").font = FONT_BOLD
    ws.cell(row=r_cap, column=7, value="Levered Beta (Relevering)").font = FONT_BOLD
    for c in range(2, 6): ws.cell(row=r_cap, column=c).border = BORDER_DOTTED
    for c in range(7, 9): ws.cell(row=r_cap, column=c).border = BORDER_DOTTED

    ws.cell(row=r_cap + 1, column=4, value="Current").font = FONT_MAIN
    ws.cell(row=r_cap + 1, column=5, value="Target (peer median)").font = FONT_MAIN

    ws.cell(row=r_cap + 2, column=2, value="Total Debt").font = FONT_MAIN
    c_cs_debt = ws.cell(row=r_cap + 2, column=3, value=f"=D{target_row}")
    apply_style(c_cs_debt, font=FONT_FORMULA, num_format='#,##0.0')
    ws.cell(row=r_cap + 2, column=4, value=f"=C{r_cap + 2}/C{r_cap + 4}").number_format = '0.00%'
    # Target weights use the peer MEDIAN, consistent with the median unlevered beta below
    ws.cell(row=r_cap + 2, column=5, value=f"=H{med_row}").number_format = '0.00%'

    ws.cell(row=r_cap + 3, column=2, value="Market Capitalization").font = FONT_MAIN
    c_cs_mcap = ws.cell(row=r_cap + 3, column=3, value=f"=E{target_row}")
    apply_style(c_cs_mcap, font=FONT_FORMULA, num_format='#,##0.0')
    ws.cell(row=r_cap + 3, column=4, value=f"=C{r_cap + 3}/C{r_cap + 4}").number_format = '0.00%'
    ws.cell(row=r_cap + 3, column=5, value=f"=1-E{r_cap + 2}").number_format = '0.00%'

    ws.cell(row=r_cap + 4, column=2, value="Total Capitalization").font = FONT_MAIN
    ws.cell(row=r_cap + 4, column=3, value=f"=C{r_cap + 2}+C{r_cap + 3}").font = FONT_FORMULA
    ws.cell(row=r_cap + 4, column=4, value=f"=D{r_cap + 2}+D{r_cap + 3}").number_format = '0.00%'
    ws.cell(row=r_cap + 4, column=5, value=f"=E{r_cap + 2}+E{r_cap + 3}").number_format = '0.00%'

    ws.cell(row=r_cap + 6, column=2, value="Debt / Equity").font = FONT_MAIN
    ws.cell(row=r_cap + 6, column=4, value=f"=C{r_cap + 2}/C{r_cap + 3}").number_format = '0.00%'
    ws.cell(row=r_cap + 6, column=5, value=f"=E{r_cap + 2}/E{r_cap + 3}").number_format = '0.00%'

    ws.cell(row=r_cap + 2, column=7, value="Peer Median Unlevered Beta").font = FONT_MAIN
    ws.cell(row=r_cap + 2, column=8, value=f"=J{med_row}").number_format = '0.00'

    ws.cell(row=r_cap + 3, column=7, value="Target Debt/ Equity").font = FONT_MAIN
    ws.cell(row=r_cap + 3, column=8, value=f"=E{r_cap + 6}").number_format = '0.00%'

    ws.cell(row=r_cap + 4, column=7, value="Tax Rate").font = FONT_MAIN
    ws.cell(row=r_cap + 4, column=8, value=f"=F{target_row}").number_format = '0.00%'

    ws.cell(row=r_cap + 5, column=7, value="Levered Beta").font = FONT_MAIN
    c_final_beta = ws.cell(row=r_cap + 5, column=8, value=f"=H{r_cap + 2}*(1+(1-H{r_cap + 4})*H{r_cap + 3})")
    apply_style(c_final_beta, font=FONT_FORMULA, num_format='0.00', border=BORDER_TB)

    # --- FINAL WACC SECTION ---
    r_final = r_cap + 8
    ws.cell(row=r_final, column=7, value="Weighted Average Cost of Capital").font = FONT_BOLD
    for c in range(7, 9): ws.cell(row=r_final, column=c).border = BORDER_DOTTED

    ws.cell(row=r_final + 2, column=7, value="Cost of Debt").font = FONT_MAIN
    ws.cell(row=r_final + 2, column=8, value=f"=D{r_cde + 4}").number_format = '0.00%'
    ws.cell(row=r_final + 3, column=7, value="Debt Weight").font = FONT_MAIN
    ws.cell(row=r_final + 3, column=8, value=f"=E{r_cap + 2}").number_format = '0.00%'

    ws.cell(row=r_final + 5, column=7, value="Cost of Equity").font = FONT_MAIN
    ws.cell(row=r_final + 5, column=8, value=f"=H{r_cde + 5}").number_format = '0.00%'
    ws.cell(row=r_final + 6, column=7, value="Equity Weight").font = FONT_MAIN
    ws.cell(row=r_final + 6, column=8, value=f"=E{r_cap + 3}").number_format = '0.00%'

    ws.cell(row=r_final + 8, column=7, value="WACC").font = FONT_BOLD
    c_wacc = ws.cell(row=r_final + 8, column=8, value=f"=(H{r_final + 2}*H{r_final + 3})+(H{r_final + 5}*H{r_final + 6})")
    apply_style(c_wacc, font=FONT_BOLD, fill=FILL_RESULT, num_format='0.00%', border=BORDER_TOP_BOTTOM_THICK)

    # --- FOOTNOTES ---
    kd_note = {
        "spread": "• Pre-Tax Cost of Debt = Risk-Free Rate + credit spread (synthetic rating from EBIT / interest expense; spread table is approximate).",
        "historical": "• Pre-Tax Cost of Debt = latest FY interest expense / total debt (a historical average coupon, not a marginal cost).",
        "manual": "• Pre-Tax Cost of Debt is a user input (rating or bond yield).",
    }[mode]
    fn_row = r_final + 4
    notes = [
        "1. Tax Rate considered as Marginal Tax Rate for the country.",
        f"2. Levered Beta is based on up to {beta_settings['years']} years of {beta_settings['freq']} data (limited by IPO date).",
        "3. Unlevered Beta = Levered Beta / (1 + (1 - Tax Rate) x Debt/Equity).",
        "4. Peer average / median exclude the target company.",
        "General Methodology:",
        "• Adjusted Levered Beta uses standard Blume Adjustment (2/3 Raw Beta + 1/3 Market Beta).",
        "• Target beta = peer median unlevered beta, relevered at the peer median debt/capital.",
        kd_note,
        "• Risk-Free Rate is based on the 10-year government bond yield.",
        "• Equity Risk Premium is a user-selected input; its source is cited on the COST OF EQUITY sheet.",
    ]
    for i, note in enumerate(notes):
        c = ws.cell(row=fn_row + i, column=2, value=note)
        apply_style(c, font=FONT_BOLD if note == "General Methodology:" else FONT_ITALIC)

    # --- SENSITIVITY: WACC vs unlevered beta and ERP ---
    s_top = fn_row + len(notes) + 2
    ws.cell(row=s_top, column=2, value="WACC Sensitivity: Peer Median Unlevered Beta (rows) vs Equity Risk Premium (columns)").font = FONT_BOLD
    for c in range(2, 9): ws.cell(row=s_top, column=c).border = BORDER_DOTTED
    hdr = s_top + 2
    ws.cell(row=hdr, column=2, value="Unlevered Beta \\ ERP").font = FONT_HEADER
    erp_deltas = [-0.02, -0.01, 0.0, 0.01, 0.02]
    beta_deltas = [-0.2, -0.1, 0.0, 0.1, 0.2]
    for j, d in enumerate(erp_deltas):
        c = ws.cell(row=hdr, column=3 + j, value=f"=$H${r_cde + 3}+({d})")
        apply_style(c, font=FONT_HEADER, num_format='0.00%', border=BORDER_BOTTOM, align=ALIGN_RIGHT)
    for i, bd in enumerate(beta_deltas):
        r = hdr + 1 + i
        cb = ws.cell(row=r, column=2, value=f"=$H${r_cap + 2}+({bd})")
        apply_style(cb, font=FONT_HEADER, num_format='0.00', align=ALIGN_RIGHT)
        for j in range(len(erp_deltas)):
            col = openpyxl.utils.get_column_letter(3 + j)
            f = (f"=$H${r_final + 2}*$H${r_final + 3}"
                 f"+($H${rf_row}+($B{r}*(1+(1-$H${r_cap + 4})*$H${r_cap + 3}))*{col}${hdr})*$H${r_final + 6}")
            cell = ws.cell(row=r, column=3 + j, value=f)
            apply_style(cell, font=FONT_FORMULA, num_format='0.00%',
                        fill=FILL_RESULT if (bd == 0.0 and erp_deltas[j] == 0.0) else None)
    ws.cell(row=hdr + 7, column=2,
            value="Centre cell equals the WACC above; each step changes beta by 0.10 and ERP by 1.0 percentage point.").font = FONT_ITALIC

    # Clean Spacing
    ws.column_dimensions['B'].width = 28
    ws.column_dimensions['G'].width = 30
    for col in ['C', 'D', 'E', 'F', 'H', 'I', 'J']:
        ws.column_dimensions[col].width = 15
    for col in ['K', 'L', 'M']:
        ws.column_dimensions[col].width = 10
    ws.column_dimensions['N'].width = 18
    ws.column_dimensions['E'].width = 20

    return {"rf_row": rf_row, "wacc_row": r_final + 8, "sens_center": (hdr + 3, 5)}


def build_coe_sheet(ws, benchmark_ticker, div_yield, erp_spec):
    """
    Cost of Equity sheet: historical index returns (reference) plus the SELECTED ERP input that the
    WACC sheet actually uses (cell G17). Returns (n_year_end_prices, n_annual_returns).
    """
    ws.title = "COST OF EQUITY"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions['A'].width = 1.15
    ws.column_dimensions['B'].width = 10
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 4
    ws.column_dimensions['F'].width = 30
    ws.column_dimensions['G'].width = 16

    headers = [
        (2, "Year", ALIGN_LEFT),
        (3, "Closing Price", ALIGN_RIGHT),
        (4, "Return", ALIGN_RIGHT)
    ]
    for col, text, align in headers:
        c = ws.cell(row=3, column=col, value=text)
        apply_style(c, font=FONT_BOLD, align=align, border=BORDER_BOTTOM)

    # --- INDEPENDENT 25-YEAR DATA PULL FOR HISTORICAL REFERENCE ---
    print(f"\n[Orchestrator] Fetching historical data for {benchmark_ticker} for Cost of Equity...")
    with contextlib.redirect_stderr(io.StringIO()):
        end_date = pd.Timestamp.today()
        start_date = end_date - pd.DateOffset(years=25)
        b_data = yf.download(benchmark_ticker, start=start_date, end=end_date, interval="1mo", progress=False)

    if isinstance(b_data.columns, pd.MultiIndex):
        top_level = b_data.columns.get_level_values(0)
        if "Adj Close" in top_level:
            b_prices = b_data["Adj Close"].iloc[:, 0]
        else:
            b_prices = b_data["Close"].iloc[:, 0]
    else:
        if "Adj Close" in b_data.columns:
            b_prices = b_data["Adj Close"]
        else:
            b_prices = b_data["Close"]

    b_prices = b_prices.dropna()
    annual_prices = b_prices.groupby(b_prices.index.year).last()

    # Drop the current partial year so it doesn't drag down the historical average
    current_year = pd.Timestamp.today().year
    if current_year in annual_prices.index:
        annual_prices = annual_prices.drop(current_year)

    n_prices = len(annual_prices)
    n_returns = max(n_prices - 1, 0)

    ws["B1"] = f"{benchmark_ticker} Historical Returns ({n_prices} year-end prices, {n_returns} annual returns)"
    apply_style(ws["B1"], font=FONT_BOLD)

    row = 4
    for i, (year, price) in enumerate(annual_prices.items()):
        c_year = ws.cell(row=row, column=2, value=int(year))
        apply_style(c_year, align=ALIGN_LEFT)

        c_price = ws.cell(row=row, column=3, value=float(price))
        apply_style(c_price, num_format='#,##0.00', align=ALIGN_RIGHT)

        if i > 0:
            c_ret = ws.cell(row=row, column=4, value=f"=(C{row}/C{row-1})-1")
            apply_style(c_ret, num_format='0.00%', align=ALIGN_RIGHT)

        row += 1

    first_row, start_row, end_row = 4, 5, row - 1

    def put(label_cell, label, value_cell, value, fmt, font=FONT_FORMULA, fill=None, bold=False):
        ws[label_cell] = label
        apply_style(ws[label_cell], font=FONT_BOLD if bold else FONT_MAIN, align=ALIGN_LEFT)
        ws[value_cell] = value
        apply_style(ws[value_cell], font=font, fill=fill, num_format=fmt, align=ALIGN_RIGHT)

    ws["F3"] = "Historical statistics (reference)"
    apply_style(ws["F3"], font=FONT_BOLD, border=BORDER_BOTTOM)
    ws["G3"].border = BORDER_BOTTOM

    if n_returns >= 2:
        put("F4", "Average returns (arithmetic)", "G4", f"=AVERAGE(D{start_row}:D{end_row})", '0.00%')
        put("F5", "Dividend yield", "G5", div_yield, '0.00%', font=FONT_INPUT, fill=FILL_INPUT)
        put("F6", "Total market return (arithmetic)", "G6", "=G4+G5", '0.00%', bold=True)
        put("F8", "Geometric mean return", "G8", f"=(C{end_row}/C{first_row})^(1/COUNT(D{start_row}:D{end_row}))-1", '0.00%')
        put("F9", "Std dev of annual returns", "G9", f"=STDEV(D{start_row}:D{end_row})", '0.00%')
        put("F10", "Std error of the mean", "G10", "=G9/SQRT(G11)", '0.00%')
        put("F11", "Number of annual returns", "G11", f"=COUNT(D{start_row}:D{end_row})", '0')
        put("F13", "Risk-free rate (WACC sheet)", "G13", None, '0.00%', font=FONT_LINK)
        put("F14", "Historical ERP: arithmetic", "G14", "=G6-G13", '0.00%')
        put("F15", "Historical ERP: geometric", "G15", "=G8+G5-G13", '0.00%')
    else:
        ws["F4"] = "Index history unavailable"
        apply_style(ws["F4"], font=FONT_WARN)

    # --- SELECTED ERP: the cell the WACC sheet links to ---
    ws["F17"] = "SELECTED EQUITY RISK PREMIUM"
    apply_style(ws["F17"], font=FONT_BOLD, align=ALIGN_LEFT)
    if erp_spec["mode"] == "input":
        c_sel = ws["G17"]
        c_sel.value = erp_spec["erp"]
        apply_style(c_sel, font=FONT_INPUT, fill=FILL_INPUT, num_format='0.00%', align=ALIGN_RIGHT, border=BORDER_TB)
    else:
        if n_returns < 2:
            raise RuntimeError("Historical ERP fallback selected but index history could not be downloaded; enter a sourced ERP instead.")
        c_sel = ws["G17"]
        c_sel.value = "=G15"
        apply_style(c_sel, font=FONT_FORMULA, fill=FILL_INPUT, num_format='0.00%', align=ALIGN_RIGHT, border=BORDER_TB)
    ws["F18"] = "Source"
    apply_style(ws["F18"], font=FONT_MAIN, align=ALIGN_LEFT)
    ws["G18"] = erp_spec["source"]
    apply_style(ws["G18"], font=FONT_ITALIC, align=ALIGN_LEFT)

    return n_prices, n_returns


def build_regression_sheet(ws, ticker, raw_prices, benchmark):
    """Builds individual beta regressions strictly dropping NaNs to prevent #DIV/0!"""
    ws.title = f"Beta - {ticker}"
    ws.sheet_view.showGridLines = False

    ws.column_dimensions['A'].width = 1.15
    ws.column_dimensions['B'].width = 12
    ws.column_dimensions['C'].width = 14
    ws.column_dimensions['D'].width = 10
    ws.column_dimensions['E'].width = 3
    ws.column_dimensions['F'].width = 12
    ws.column_dimensions['G'].width = 14
    ws.column_dimensions['H'].width = 10
    ws.column_dimensions['I'].width = 3
    ws.column_dimensions['J'].width = 18
    ws.column_dimensions['K'].width = 12

    ws["A1"] = f"Regression Beta - {ticker}"
    apply_style(ws["A1"], font=FONT_BOLD)

    headers = [
        (2, f"{ticker} Returns", FILL_HEADER),
        (6, f"{benchmark} Returns", FILL_HEADER),
        (10, "Beta Drifting", FILL_HEADER)
    ]
    for col, text, fill in headers:
        c = ws.cell(row=3, column=col, value=text)
        apply_style(c, font=FONT_HEADER, fill=fill, align=ALIGN_CENTER)
        if col != 10:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+2)
        else:
            ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col+1)

    sub_headers = [
        (2, "Date", ALIGN_LEFT), (3, "Closing Price", ALIGN_RIGHT), (4, "Return", ALIGN_RIGHT),
        (6, "Date", ALIGN_LEFT), (7, "Closing Price", ALIGN_RIGHT), (8, "Return", ALIGN_RIGHT)
    ]
    for col, text, align in sub_headers:
        c = ws.cell(row=5, column=col, value=text)
        apply_style(c, font=FONT_BOLD, align=align, border=BORDER_BOTTOM)

    # Defensive check to prevent KeyError if yfinance drops columns
    if ticker in raw_prices.columns and benchmark in raw_prices.columns:
        valid_data = raw_prices[[ticker, benchmark]].dropna()
    else:
        # Prevents Python crash, allows Excel generator to output blank layout
        valid_data = pd.DataFrame(columns=[ticker, benchmark])

    row = 6
    for i, (date, row_data) in enumerate(valid_data.iterrows()):
        c_date = ws.cell(row=row, column=2, value=date.strftime('%Y-%m-%d'))
        apply_style(c_date, align=ALIGN_LEFT)

        c_price = ws.cell(row=row, column=3, value=row_data[ticker])
        apply_style(c_price, num_format='#,##0.00', align=ALIGN_RIGHT)

        if i > 0:
            c_ret = ws.cell(row=row, column=4, value=f"=(C{row}/C{row-1})-1")
            apply_style(c_ret, num_format='0.00%', align=ALIGN_RIGHT)

        c_bdate = ws.cell(row=row, column=6, value=date.strftime('%Y-%m-%d'))
        apply_style(c_bdate, align=ALIGN_LEFT)

        c_bprice = ws.cell(row=row, column=7, value=row_data[benchmark])
        apply_style(c_bprice, num_format='#,##0.00', align=ALIGN_RIGHT)

        if i > 0:
            c_bret = ws.cell(row=row, column=8, value=f"=(G{row}/G{row-1})-1")
            apply_style(c_bret, num_format='0.00%', align=ALIGN_RIGHT)

        row += 1

    end_row = row - 1 if row > 6 else 6

    ws["J5"] = "Levered Raw Beta"
    ws.cell(row=5, column=11, value=f"=SLOPE(D7:D{end_row}, H7:H{end_row})").number_format = '0.00'

    ws["J6"] = "Raw Beta Weight"
    apply_style(ws.cell(row=6, column=11, value=(2/3)), font=FONT_INPUT, num_format='0.00%')

    ws["J8"] = "Market Beta"
    apply_style(ws.cell(row=8, column=11, value=1.00), font=FONT_INPUT, num_format='0.00')

    ws["J9"] = "Market Beta Weight"
    apply_style(ws.cell(row=9, column=11, value=(1/3)), font=FONT_INPUT, num_format='0.00%')

    ws["J11"] = "Adjusted Beta"
    apply_style(ws["J11"], font=FONT_BOLD)
    c_adj = ws.cell(row=11, column=11, value="=(K5*K6)+(K8*K9)")
    apply_style(c_adj, font=FONT_BOLD, fill=PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid"), num_format='0.00')

    ws["J15"] = "SUMMARY OUTPUT"
    ws["J16"] = "Regression Statistics"
    apply_style(ws["J16"], font=FONT_ITALIC, border=BORDER_BOTTOM)
    ws.merge_cells("J16:K16")

    stats = [
        ("Multiple R", f"=CORREL(D7:D{end_row}, H7:H{end_row})", '0.0000'),
        ("R Square", f"=RSQ(D7:D{end_row}, H7:H{end_row})", '0.0000'),
        ("Standard Error", f"=STEYX(D7:D{end_row}, H7:H{end_row})", '0.0000'),
        ("Observations", f"=COUNT(D7:D{end_row})", '0'),
        ("Beta Std Error", f"=K19/SQRT(DEVSQ(H7:H{end_row}))", '0.0000'),
        ("Raw Beta 95% CI Low", "=K5-TINV(0.05,K20-2)*K21", '0.00'),
        ("Raw Beta 95% CI High", "=K5+TINV(0.05,K20-2)*K21", '0.00'),
    ]

    r = 17
    for label, formula, fmt in stats:
        ws.cell(row=r, column=10, value=label)
        ws.cell(row=r, column=11, value=formula).number_format = fmt
        r += 1


def main():
    print("=" * 50)
    print("      WACC MODEL EXCEL GENERATOR (FINAL)      ")
    print("=" * 50)

    # 1. Gather Inputs & Data
    macro = get_macro_inputs()
    target_ticker = macro['ticker']

    kd_spec = get_cost_of_debt_spec(macro)

    print("\n[Orchestrator] Handing off to Beta Regression Engine...")
    beta_results = run_beta_comps(macro)

    # --- DYNAMIC DIVIDEND YIELD FETCH (used only for the historical reference block) ---
    bench_ticker = beta_results['benchmark']
    bench_info = yf.Ticker(bench_ticker).info or {}
    div_yield = bench_info.get('trailingAnnualDividendYield') or bench_info.get('dividendYield')

    if div_yield is None:
        print(f"\n[!] Could not fetch Dividend Yield for {bench_ticker} automatically (common for indices).")
        div_yield = prompt_float(f"Enter Dividend Yield % for {bench_ticker} [Default 1.30%]: ", 1.30, "dividend yield")
    else:
        print(f"\n✓ Fetched Dividend Yield for {bench_ticker}: {div_yield * 100:.2f}%")

    erp_spec = get_erp_spec()

    # 2. Setup Excel Workbook
    wb = openpyxl.Workbook()
    wb.calculation.fullCalcOnLoad = True  # make Excel compute every formula on open

    print("\n[Orchestrator] Building Excel Model Tabs...")

    beta_settings = {
        "years": beta_results['period_years'],
        "freq": "monthly" if beta_results['frequency'] == "1mo" else "weekly" if beta_results['frequency'] == "1wk" else "daily"
    }

    ws_wacc = wb.active
    ws_coe = wb.create_sheet("COST OF EQUITY")
    build_coe_sheet(ws_coe, bench_ticker, div_yield, erp_spec)

    comps = beta_results['summary_table'].to_dict('records')
    refs = build_wacc_sheet(ws_wacc, macro, comps, beta_settings, kd_spec, erp_spec)

    # Historical reference block on the CoE sheet needs the Rf cell that lives on the WACC sheet
    if ws_coe["F13"].value:
        ws_coe["G13"] = f"=WACC!$H${refs['rf_row']}"

    # Sheets 3+: Beta Regressions
    raw_prices = beta_results['raw_prices']
    for row_data in comps:
        ticker = row_data['Ticker']
        local_bench = row_data.get('Benchmark', bench_ticker)

        ws_beta = wb.create_sheet(f"Beta - {ticker}")
        build_regression_sheet(ws_beta, ticker, raw_prices, local_bench)

    # 3. Save File
    if not os.path.exists("output"): os.makedirs("output")
    filename = f"output/WACC_Model_{target_ticker}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    wb.save(filename)

    print(f"\n✓ SUCCESS! Institutional WACC model generated: {filename}")
    print("  Open in Excel and press Save once before sharing, so previewers show calculated values.")


if __name__ == "__main__":
    main()